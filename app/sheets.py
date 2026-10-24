# -*- coding: utf-8 -*-
"""读取已经做好的《XX部门空课表》，还原成成员占用表，直接用来排值班表。"""
import os, re
import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock

from core import Member

BLOCK_ROWS = [4, 8, 12, 16]
BLACK = ("00000000", "FF000000", "000000")


def _is_special(color):
    """红/蓝/绿等彩色才算部长标记；黑色、深灰、白色都算正文。"""
    if color is None:
        return False
    rgb = getattr(color, "rgb", None)
    if not isinstance(rgb, str):
        return False
    s = rgb.upper()
    if s in BLACK:
        return False
    if len(s) == 8:
        s = s[2:]
    if len(s) != 6:
        return False
    try:
        r, g, b = int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16)
    except ValueError:
        return False
    if max(r, g, b) - min(r, g, b) <= 40:   # 灰阶（黑/深灰/白）不算标记
        return False
    return True


def _weeks_from_text(s):
    weeks = set()
    for part in re.split(r"[，,、]", s or ""):
        part = part.strip()
        m = re.fullmatch(r"(\d+)\s*-\s*(\d+)", part)
        if m:
            weeks.update(range(int(m.group(1)), int(m.group(2)) + 1))
        elif re.fullmatch(r"\d+", part):
            weeks.add(int(part))
    return weeks


def _cell_lines(cell):
    """把单元格拆成 [(文本行, 是否红色)]"""
    v = cell.value
    if v is None:
        return []
    if isinstance(v, CellRichText):
        lines, cur, cur_red = [], "", None
        for block in v:
            if isinstance(block, TextBlock):
                txt = block.text
                color = block.font.color if (block.font and block.font.color) else None
                red = _is_special(color)
                for ch in txt:
                    if ch == "\n":
                        if cur.strip():
                            lines.append((cur.strip(), bool(cur_red)))
                        cur, cur_red = "", None
                    else:
                        if cur_red is None:
                            cur_red = red
                        cur += ch
        if cur.strip():
            lines.append((cur.strip(), bool(cur_red)))
        return lines
    return [(l.strip(), False) for l in str(v).split("\n") if l.strip()]


def _dept_name(ws, path):
    title = str(ws["B2"].value or "").strip()
    name = re.sub(r"(部门)?空课表$", "", title).strip()
    if not name:
        name = re.sub(r"(部门)?空课表$", "", os.path.splitext(os.path.basename(path))[0]).strip()
    if not name:
        return "未命名部门"
    try:
        from pipeline import canon_dept
        return canon_dept(name)
    except Exception:
        return name


def parse_empty_sheet(path):
    wb = openpyxl.load_workbook(path, rich_text=True)
    ws = wb.worksheets[0]
    dept = _dept_name(ws, path)

    free = {}              # name -> {(day, blk): set(weeks)}
    grid_minister = set()  # 大格里标了颜色的人
    info_minister = set()  # 底部信息表里标了颜色的人
    order = []

    for blk, row in enumerate(BLOCK_ROWS):
        for day in range(1, 6):
            cell = ws.cell(row=row, column=3 + day)
            for text, red in _cell_lines(cell):
                m = re.match(r"(.+?)[（(]([^）)]*)[）)]", text)
                if not m:
                    continue
                name = m.group(1).strip()
                weeks = _weeks_from_text(m.group(2))
                if not name or not weeks:
                    continue
                if name not in free:
                    free[name] = {}
                    order.append(name)
                free[name][(day, blk)] = weeks
                if red:
                    grid_minister.add(name)

    # 底部信息收集表
    info = {}
    for r in range(21, ws.max_row + 1):
        name = ws.cell(row=r, column=2).value
        if not name or str(name).strip() in ("姓名",):
            continue
        name = str(name).strip()
        klass = ws.cell(row=r, column=3).value or ""
        phone = ws.cell(row=r, column=4).value or ""
        c = ws.cell(row=r, column=2)
        is_red = _is_special(c.font.color if c.font else None)
        info[name] = (str(klass).strip(), str(phone).strip())
        if is_red:
            info_minister.add(name)
        if name not in free:
            free[name] = {}
            order.append(name)

    # 底部信息表的颜色更可靠（老版本的表格大格里可能整格都是红的），
    # 所以信息表里出现了颜色标记就只认信息表。
    minister = info_minister if info_minister else grid_minister

    members = []
    for name in order:
        klass, phone = info.get(name, ("", ""))
        m = Member(name, klass, phone, dept, name in minister, path)
        # 表里列的是"有空"的周次，这里反推成"被占用"的周次
        busy = {}
        for (day, blk), weeks in free[name].items():
            busy[(day, blk)] = set(range(1, 31)) - weeks
        m.busy = busy
        m.error = None
        members.append(m)
    return dept, members


def load_sheets(src_dir, order=None):
    data = {}
    files = [f for f in os.listdir(src_dir)
             if f.lower().endswith((".xlsx", ".xlsm")) and "值班表" not in f]
    seq = {b: i for i, b in enumerate(order or [])}
    files.sort(key=lambda f: (seq.get(f, 9999), f))
    for f in files:
        dept, members = parse_empty_sheet(os.path.join(src_dir, f))
        data[dept] = members
    return data
