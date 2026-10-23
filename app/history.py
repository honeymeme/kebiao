# -*- coding: utf-8 -*-
"""读取已经排好的《第X周系办值班表》，把已排的班次记下来，后面几周接着排。"""
import os, re
import openpyxl

BLOCK_ROWS = [3, 7, 11, 15]
CN = {"零": 0, "一": 1, "二": 2, "三": 3, "四": 4, "五": 5,
      "六": 6, "七": 7, "八": 8, "九": 9}


def cn_to_int(s):
    s = (s or "").strip()
    if not s:
        return None
    if s.isdigit():
        return int(s)
    if s == "十":
        return 10
    if "十" in s:
        a, _, b = s.partition("十")
        tens = CN.get(a, 1) if a else 1
        ones = CN.get(b, 0) if b else 0
        return tens * 10 + ones
    return CN.get(s)


def _week_of(path, ws):
    title = str(ws["A1"].value or "")
    m = re.search(r"第\s*([零一二三四五六七八九十\d]+)\s*周", title)
    if m:
        return cn_to_int(m.group(1))
    m = re.search(r"第\s*([零一二三四五六七八九十\d]+)\s*周", os.path.basename(path))
    return cn_to_int(m.group(1)) if m else None


def parse_duty_sheet(path):
    """返回 (第几周, {(blk, day): [姓名...]})"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    week = _week_of(path, ws)
    slots = {}
    for blk, row in enumerate(BLOCK_ROWS):
        for day in range(1, 6):
            v = ws.cell(row=row, column=2 + day).value
            if not v:
                continue
            text = str(v)
            tokens = [t.strip() for t in re.split(r"[\n/／、,，;；]+", text) if t.strip()]
            names = []
            i = 0
            while i < len(tokens):
                t = tokens[i]
                if re.fullmatch(r"1\d{10}", t):
                    i += 1
                    continue
                if (re.fullmatch(r"[\u4e00-\u9fa5·]{2,6}", t)
                        and i + 1 < len(tokens)
                        and re.fullmatch(r"1\d{10}", tokens[i + 1])):
                    names.append(t)
                    i += 2
                else:
                    i += 1
            if names:
                slots[(blk, day)] = names
    return week, slots


def load_history(src_dir):
    """扫描目录里所有"值班表"xlsx，返回 (每个名字已排次数, 已排周次集合, 明细)"""
    counts = {}
    weeks = set()
    detail = {}
    for f in sorted(os.listdir(src_dir)):
        if not f.lower().endswith((".xlsx", ".xlsm")):
            continue
        if "值班表" not in f:
            continue
        try:
            week, slots = parse_duty_sheet(os.path.join(src_dir, f))
        except Exception:
            continue
        if week is None:
            continue
        weeks.add(week)
        detail[week] = slots
        for names in slots.values():
            for n in names:
                counts[n] = counts.get(n, 0) + 1
    return counts, weeks, detail
